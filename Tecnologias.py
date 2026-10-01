from Wappalyzer import Wappalyzer
import argparse
from openpyxl import Workbook
import logging
import shutil 



logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s:%(message)s")
wb = Workbook() 
ws = wb.active
b=1
def tecnologi(url):
    Wappalyzer = Wappalyzer.latest(update=True)
    try:
        #warnings.simplefilter("ignore")
        WebPage = WebPage.new_from_url("https://www.example.com")
        tecnologias = Wappalyzer.analyze(WebPage)
        for t in tecnologias:
            A = str("Tecnologias detectadas: {}".format(t))
            ws['A' + str(b)] = A
            b+=1
    except:
        A = str("Ha ocurrido un error intente de nuevo")
        ws['A' + str(b)] = A
        b+=1

wb.save('Tecnologias.xlsx')

shutil.move("tecnologias.xlxs","reporte")
